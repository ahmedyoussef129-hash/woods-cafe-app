from pathlib import Path
import sqlite3
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
XLSX=ROOT.parent/"woods cafe.xlsx"
DB=ROOT/"woods_cafe.db"

if not XLSX.exists():
    raise SystemExit(f"Excel not found: {XLSX}")

wb=load_workbook(XLSX, read_only=True, data_only=True)
ws=wb["المبيعات"]
products={}
# Based on the verified workbook structure: B=product, M=base price.
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    name=row[1] if len(row)>1 else None
    price=row[12] if len(row)>12 else None
    if isinstance(name,str) and name.strip():
        try: price=float(price or 0)
        except: price=0
        products[name.strip()]=price

c=sqlite3.connect(DB)
for name,price in products.items():
    c.execute("INSERT OR IGNORE INTO products(name,base_price) VALUES(?,?)",(name,price))
c.commit(); c.close()
print(f"Imported {len(products)} products.")
